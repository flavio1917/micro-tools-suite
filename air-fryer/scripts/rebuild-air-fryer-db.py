#!/usr/bin/env python3
import argparse
import json
import re
import unicodedata
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

LANGUAGE_MARKERS = {
    "it": (" il ", " la ", " di ", " non ", " contatta ", " assistenza ", " friggitrice ", " scollega ", " dispositivo "),
    "en": (" the ", " there is ", " contact ", " customer support ", " unplug ", " air fryer ", " has been "),
}
MARKER_RE = re.compile(
    r'^\s*(?:il\s+display\s+mostra\s+il\s+codice\s+errore\s*)?["“”]?\s*'
    r'(E\d{1,2}(?:\s*(?:/|–|-)\s*E?\d{1,2})*)\s*["“”]?\s*[:.]?\s*',
    re.IGNORECASE | re.MULTILINE,
)


def clean(value):
    if value is None:
        return ""
    return str(value).replace("\r\n", "\n").replace("\r", "\n").strip()


def slugify(value):
    value = unicodedata.normalize("NFKD", clean(value)).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return value


def detect_language(text):
    lowered = f" {clean(text).lower()} "
    scores = {lang: sum(lowered.count(marker) for marker in markers) for lang, markers in LANGUAGE_MARKERS.items()}
    return "en" if scores["en"] > scores["it"] else "it"


def normalize_code(code):
    return re.sub(r"\s+", "", code.upper().replace("–", "-"))


def expand_code_group(group):
    normalized = normalize_code(group)
    parts = re.split(r"[/\-]", normalized)
    output = []
    for part in parts:
        if not part:
            continue
        if not part.startswith("E"):
            part = f"E{part}"
        if part not in output:
            output.append(part)
    return output


def classify_error(code, text):
    lowered = clean(text).lower()
    support_terms = ("assistenza", "support", "customer service", "contact customer", "contattare")
    reset_terms = ("raffreddare", "cool completely", "lasciarla raffreddare", "reinserisci", "reinsert")
    voltage_terms = ("220", "240v", "120v", "presa", "outlet")
    divider_terms = ("divisorio", "divider")
    if any(term in lowered for term in divider_terms):
        return "caution", True, "caution"
    if any(term in lowered for term in reset_terms) and not any(term in lowered for term in support_terms):
        return "software_reset", True, "verifica_utente"
    if any(term in lowered for term in voltage_terms):
        return "caution", True, "fermare_e_verificare"
    if any(term in lowered for term in support_terms):
        return "hardware", False, "assistenza"
    return "caution", False, "fermare_e_verificare"


def split_error_blocks(raw):
    raw = clean(raw)
    matches = list(MARKER_RE.finditer(raw))
    if not matches:
        return []
    blocks = []
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(raw)
        text = raw[match.end():end].strip(" \n\t.-")
        if not text:
            continue
        group = normalize_code(match.group(1))
        codes = expand_code_group(group)
        severity, diy_fixable, action_level = classify_error(group, text)
        for code in codes:
            blocks.append({
                "code": code,
                "display_code": group,
                "aliases": [candidate for candidate in codes if candidate != code],
                "source_text": text,
                "source_language": detect_language(text),
                "severity": severity,
                "diy_fixable": diy_fixable,
                "action_level": action_level,
                "evidence_level": "model_specific",
            })
    return blocks


def parse_specs(raw):
    text = clean(raw)
    specs = {}
    capacity = re.search(r"capacit[àa]\s*:\s*([\d.,]+)\s*(?:l|litri|qt)", text, re.IGNORECASE)
    watts = re.search(r"(?:assorbimento|rated power)\s*:\s*([\d.,]+)\s*w", text, re.IGNORECASE)
    baskets = re.search(r"numero di cestelli\s*:\s*(\d+)", text, re.IGNORECASE)
    if capacity:
        specs["capacity_liters"] = float(capacity.group(1).replace(",", "."))
    if watts:
        specs["wattage"] = float(watts.group(1).replace(",", "."))
    if baskets:
        specs["basket_count"] = int(baskets.group(1))
    return specs


def cell(row, headers, name):
    return clean(row[headers[name]]) if name in headers else ""


def build_model(row, headers):
    brand = cell(row, headers, "Marca")
    model_name = cell(row, headers, "Modello")
    model_spec = cell(row, headers, "Model Spec")
    raw_errors = cell(row, headers, "Errori")
    manual_url = cell(row, headers, "Link Manuale Utente")
    specs_url = cell(row, headers, "Link Specifiche")
    canonical_slug = f"{slugify(brand)}-{slugify(model_spec or model_name)}"
    errors = split_error_blocks(raw_errors)
    source = {"manual_url": manual_url, "specs_url": specs_url}
    return {
        "canonical_slug": canonical_slug,
        "route_slug": slugify(model_name),
        "brand_slug": slugify(brand),
        "brand": brand,
        "primary_display_name": model_name,
        "display_names": [model_name],
        "model_spec": model_spec,
        "specs": parse_specs(cell(row, headers, "Specifiche")),
        "error_codes": [{**error, "source": source} for error in errors],
        "symptoms": [],
        "source": source,
        "_raw_errors": raw_errors,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }


def main():
    parser = argparse.ArgumentParser(description="Build the air-fryer runtime database from Errori.xlsx.")
    parser.add_argument("--input", default="../../Errori.xlsx")
    parser.add_argument("--output", default="src/data/air-fryer-db.json")
    parser.add_argument("--report", default="scripts/air-fryer-import-report.json")
    args = parser.parse_args()

    workbook = load_workbook(args.input, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise SystemExit("The workbook is empty.")
    headers = {clean(name): index for index, name in enumerate(rows[0]) if clean(name)}
    required = {"Marca", "Modello", "Model Spec", "Errori"}
    missing = required - set(headers)
    if missing:
        raise SystemExit(f"Missing required columns: {', '.join(sorted(missing))}")

    models = [build_model(row, headers) for row in rows[1:] if cell(row, headers, "Marca") and cell(row, headers, "Modello")]
    duplicate_slugs = [slug for slug, count in Counter(model["canonical_slug"] for model in models).items() if count > 1]
    report_rows = [
        {
            "brand": model["brand"],
            "model": model["primary_display_name"],
            "model_spec": model["model_spec"],
            "canonical_slug": model["canonical_slug"],
            "code": error["code"],
            "display_code": error["display_code"],
            "source_language": error["source_language"],
            "source_text_length": len(error["source_text"]),
            "manual_url": error["source"]["manual_url"],
        }
        for model in models for error in model["error_codes"]
    ]
    empty_blocks = [item for item in report_rows if item["source_text_length"] == 0]
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "input": str(Path(args.input)),
        "model_count": len(models),
        "error_code_count": len(report_rows),
        "source_languages": dict(Counter(item["source_language"] for item in report_rows)),
        "duplicate_canonical_slugs": duplicate_slugs,
        "empty_error_blocks": empty_blocks,
        "rows": report_rows,
    }

    output_path = Path(args.output)
    report_path = Path(args.report)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(models, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"Generated {output_path}: {len(models)} models, {len(report_rows)} error-code records.")
    print(f"Generated {report_path}.")
    if duplicate_slugs:
        print(f"Warning: duplicate canonical slugs: {', '.join(duplicate_slugs)}")
    if empty_blocks:
        raise SystemExit("Import failed: one or more error blocks are empty.")


if __name__ == "__main__":
    main()